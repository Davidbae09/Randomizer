import tkinter
from tkinter import filedialog, ttk, messagebox, simpledialog
import tkinter.font as tkFont
from PIL import Image, ImageTk
import openpyxl  # For handling Excel files
from docx import Document  # For handling Word files
from modul import (
    TournamentBracketGeneratorFrame, RandomItemPickerFrame, GroupGeneratorFrame, 
    ColorGeneratorFrame, RandomCardGeneratorFrame, RandomDiceGeneratorFrame, 
    RandomCoinGeneratorFrame, RandomNumberSetGeneratorFrame, RandomMazeGeneratorFrame, 
    RandomColorPickerFrame, RandomPasswordGeneratorFrame, RandomCoordinateGeneratorFrame
)
import os
from translation import translate, load_language_setting, CONFIG_FILE

class MainApp:
    def __init__(self, master):
        self.master = master
        master.title("RANDOMIZER")
        
        try:
            self.master.iconbitmap('randomizer.ico')
        except Exception as e:
            print(f"Icon error: {e}")

        # Load language setting
        self.language = load_language_setting()
        
        #changing font
        #self.custom_font = tkFont.Font(family="Helvetica", size=12, weight="bold")
        #master.option_add("*Font", self.custom_font)
        #self.background_color = 'lightblue'  # Define your background color once
        
        #master.configure(bg=self.background_color)  # Set main window background color

        # Maximize the window
        master.state('zoomed')

        # Create menu bar
        self.menu_bar = tkinter.Menu(master)
        master.config(menu=self.menu_bar)

        # Menu options
        self.menu = tkinter.Menu(self.menu_bar, tearoff=0)
        self.menu_bar.add_cascade(label=translate("Menu"), menu=self.menu)

        # Add Random under Menu
        self.menu.add_command(label=translate("Random Item Picker"), command=self.show_random_item_picker)
        self.menu.add_command(label=translate("Group Generator"), command=self.show_group_generator)
        self.menu.add_command(label=translate("Tournament Bracket Generator"), command=self.show_tournament_bracket_generator)
        self.menu.add_command(label=translate("Coloring Generator"), command=self.show_color_generator)
        self.menu.add_command(label=translate("Card Generator"), command=self.show_card_generator)
        self.menu.add_command(label=translate("Dice Generator"), command=self.show_dice_generator)
        self.menu.add_command(label=translate("Coin Generator"), command=self.show_coin_generator)
        self.menu.add_command(label=translate("Random Number Generator"), command=self.show_number_set_generator)
        self.menu.add_command(label=translate("Random Maze Generator"), command=self.show_maze_generator)
        self.menu.add_command(label=translate("Random Color Generator"), command=self.show_random_color_generator)
        self.menu.add_command(label=translate("Random Password Generator"), command=self.show_password_generator)
        self.menu.add_command(label=translate("Random Coordinate Generator"), command=self.show_coordinate_generator)

        # Tools menu
        self.tools_menu = tkinter.Menu(self.menu_bar, tearoff=0)
        self.menu_bar.add_cascade(label=translate("Tools"), menu=self.tools_menu)

        # Add Import option under Tools menu
        self.tools_menu.add_command(label=translate("Import"), command=self.open_import)
        self.tools_menu.add_command(label=translate("Change Language"), command=self.change_language)

        # Help menu
        self.help_menu = tkinter.Menu(self.menu_bar, tearoff=0)
        self.menu_bar.add_cascade(label=translate("Help"), menu=self.help_menu)
        self.help_menu.add_command(label=translate("User Guide"), command=self.show_user_guide)
        self.help_menu.add_command(label=translate("FAQ"), command=self.show_faq)
        self.help_menu.add_command(label=translate("About"), command=self.show_about_dialog)

        # Placeholder for current tool frame
        self.current_tool_frame = None

        # Initially show the Random Item Picker
        self.show_random_item_picker()

    def change_language(self):
        self.language = 'bahasa' if self.language == 'english' else 'english'
        with open(CONFIG_FILE, 'w') as file:
            file.write(f'language={self.language}')
        self.master.destroy()
        self.master = tkinter.Tk()
        self.__init__(self.master)
        self.master.mainloop()

    def show_tournament_bracket_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = TournamentBracketGeneratorFrame(self.master)

    def show_random_item_picker(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomItemPickerFrame(self.master)

    def show_group_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = GroupGeneratorFrame(self.master)

    def show_color_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = ColorGeneratorFrame(self.master)

    def show_card_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomCardGeneratorFrame(self.master)

    def show_dice_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomDiceGeneratorFrame(self.master)

    def show_number_set_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomNumberSetGeneratorFrame(self.master)

    def show_maze_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomMazeGeneratorFrame(self.master)

    def show_coin_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomCoinGeneratorFrame(self.master)

    def show_random_color_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomColorPickerFrame(self.master)

    def show_password_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomPasswordGeneratorFrame(self.master)

    def show_coordinate_generator(self):
        if self.current_tool_frame:
            self.current_tool_frame.destroy()
        self.current_tool_frame = RandomCoordinateGeneratorFrame(self.master)

    def open_import(self):
        file_path = filedialog.askopenfilename(
            filetypes=[("All Supported Files", "*.txt;*.xlsx;*.docx"), ("Text Files", "*.txt"), ("Excel Files", "*.xlsx"), ("Word Files", "*.docx")]
        )
        if file_path:
            if file_path.endswith('.txt'):
                self.import_from_txt(file_path)
            elif file_path.endswith('.xlsx'):
                self.import_from_excel(file_path)
            elif file_path.endswith('.docx'):
                self.import_from_word(file_path)

    def import_from_txt(self, file_path):
        with open(file_path, 'r') as file:
            content = file.read()
        self.insert_content_into_current_tool(content)

    def import_from_excel(self, file_path):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        content = '\n'.join([str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value])
        self.insert_content_into_current_tool(content)

    def import_from_word(self, file_path):
        doc = Document(file_path)
        content = '\n'.join([para.text for para in doc.paragraphs])
        self.insert_content_into_current_tool(content)

    def insert_content_into_current_tool(self, content):
        if isinstance(self.current_tool_frame, RandomItemPickerFrame):
            self.current_tool_frame.item_text.delete('1.0', tkinter.END)
            self.current_tool_frame.item_text.insert(tkinter.END, content)
        elif isinstance(self.current_tool_frame, GroupGeneratorFrame):
            self.current_tool_frame.names_text.delete('1.0', tkinter.END)
            self.current_tool_frame.names_text.insert(tkinter.END, content)
        elif isinstance(self.current_tool_frame, TournamentBracketGeneratorFrame):
            self.current_tool_frame.team_text.delete('1.0', tkinter.END)
            self.current_tool_frame.team_text.insert(tkinter.END, content)
                
    def show_user_guide(self):
        user_guide_window = tkinter.Toplevel(self.master)
        user_guide_window.title(translate("User Guide"))

        # Window size and position
        window_width = 600
        window_height = 400
        screen_width = self.master.winfo_screenwidth()
        screen_height = self.master.winfo_screenheight()
        position_x = (screen_width // 2) - (window_width // 2)
        position_y = (screen_height // 2) - (window_height // 2)
        user_guide_window.geometry(f"{window_width}x{window_height}+{position_x}+{position_y}")

        # Get the default system font
        default_font = tkFont.nametofont("TkDefaultFont")

        # Frame for Text and Scrollbar
        frame = tkinter.Frame(user_guide_window)
        frame.pack(fill=tkinter.BOTH, expand=True)

        # Text widget with system default font
        text_widget = tkinter.Text(frame, wrap="word", state=tkinter.NORMAL, font=default_font)
        text_widget.pack(side=tkinter.LEFT, fill=tkinter.BOTH, expand=True)

        scrollbar = tkinter.Scrollbar(frame, command=text_widget.yview)
        scrollbar.pack(side=tkinter.RIGHT, fill=tkinter.Y)
        text_widget.config(yscrollcommand=scrollbar.set)

        # Display the user guide text
        user_guide_text = translate("guide_text")
        text_widget.insert(tkinter.END, user_guide_text)
        text_widget.config(state=tkinter.DISABLED)

        close_button = tkinter.Button(user_guide_window, text="Close", command=user_guide_window.destroy, font=default_font)
        close_button.pack(pady=10)

        user_guide_window.grab_set()
        user_guide_window.focus_set()
        user_guide_window.transient(self.master)
    def show_faq(self):
        simpledialog.messagebox.showinfo(translate("FAQ"), translate("faq_text"))
    def show_about_dialog(self):
        tkinter.messagebox.showinfo(translate("About"), "Randomizer Tool\nVersion 1.0\n@davidbae09")

    def remove_current_tool_frame(self):
        if self.current_tool_frame is not None:
            self.current_tool_frame.destroy()
            self.current_tool_frame = None

# Sample usage
if __name__ == "__main__":
    root = tkinter.Tk()
    app = MainApp(root)
    root.mainloop()
